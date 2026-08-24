"""Create deterministic journal submission assets from reported paper content."""

from pathlib import Path

import matplotlib.pyplot as plt
from matplotlib.patches import FancyArrowPatch, FancyBboxPatch
from docx import Document
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "submission_assets"
TITLE = (
    "Does crop-host supervision improve Swin-based multi-crop condition "
    "recognition? A controlled 90-run in-domain audit of performance, "
    "consistency, and confidence"
)
EN_DASH = chr(0x2013)

INK = "#18323A"
MUTED = "#52666D"
GREEN = "#2F7D59"
GREEN_LIGHT = "#E7F2EC"
BLUE = "#315B7D"
BLUE_LIGHT = "#E8EFF5"
AMBER = "#A56A1F"
AMBER_LIGHT = "#F8EEDC"
RED = "#A94843"
RED_LIGHT = "#F7E8E6"
LINE = "#B8C4C7"
WHITE = "#FFFFFF"

HIGHLIGHTS = (
    "Family-balanced testing preserves the near-null 43-class effect",
    "Ninety Swin-S runs audit host supervision across three benchmarks",
    "Flat and dual models are equivalent on all 43-class primary outcomes",
    "Ground-truth-host oracle gains 2.2 points on the hardest benchmark",
    "Temperature scaling cuts hardest-set ECE from 5.4% to about 1.6%",
)


def rounded_box(ax, x, y, width, height, face, edge=LINE, radius=0.014):
    patch = FancyBboxPatch(
        (x, y),
        width,
        height,
        boxstyle=f"round,pad=0.008,rounding_size={radius}",
        linewidth=1.25,
        edgecolor=edge,
        facecolor=face,
    )
    ax.add_patch(patch)
    return patch


def stage_header(ax, x, y, number, title, color):
    ax.text(
        x,
        y,
        number,
        ha="center",
        va="center",
        fontsize=9.5,
        fontweight="bold",
        color=WHITE,
        bbox={"boxstyle": "circle,pad=0.35", "facecolor": color, "edgecolor": color},
    )
    ax.text(
        x + 0.021,
        y,
        title,
        ha="left",
        va="center",
        fontsize=9.4,
        fontweight="bold",
        color=INK,
    )


def arrow(ax, x1, x2, y):
    ax.add_patch(
        FancyArrowPatch(
            (x1, y),
            (x2, y),
            arrowstyle="-|>",
            mutation_scale=19,
            linewidth=2.0,
            color=MUTED,
        )
    )


def make_graphical_abstract() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    fig = plt.figure(figsize=(8.0, 3.2), dpi=300, facecolor=WHITE)
    ax = fig.add_axes((0, 0, 1, 1))
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    ax.axis("off")

    ax.text(
        0.5,
        0.955,
        "Does crop-host supervision improve Swin-based multi-crop condition recognition?",
        ha="center",
        va="top",
        fontsize=11.8,
        fontweight="bold",
        color=INK,
    )
    ax.text(
        0.5,
        0.902,
        "A controlled 90-run in-domain audit of performance, consistency, and confidence",
        ha="center",
        va="top",
        fontsize=9.0,
        color=MUTED,
    )

    panel_y, panel_h, panel_w = 0.46, 0.37, 0.285
    xs = (0.035, 0.3575, 0.68)
    colors = ((GREEN_LIGHT, GREEN), (BLUE_LIGHT, BLUE), (AMBER_LIGHT, AMBER))
    for x, (face, edge) in zip(xs, colors):
        rounded_box(ax, x, panel_y, panel_w, panel_h, face, edge)

    stage_header(ax, xs[0] + 0.035, 0.785, "1", "Identity-controlled data", GREEN)
    datasets = (
        f"NPD{EN_DASH}Rice{EN_DASH}COCO{EN_DASH}43",
        f"MultiCrop{EN_DASH}88",
        f"Agri{EN_DASH}Foundation{EN_DASH}145k",
    )
    for index, (name, count) in enumerate(zip(datasets, ("43 classes", "88 classes", "215 classes"))):
        y = 0.706 - 0.068 * index
        ax.text(xs[0] + 0.022, y, name, fontsize=7.5, fontweight="bold", color=INK, va="center")
        ax.text(xs[0] + panel_w - 0.02, y, count, fontsize=7.2, color=GREEN, va="center", ha="right")
    ax.plot([xs[0] + 0.02, xs[0] + panel_w - 0.02], [0.542, 0.542], color=LINE, lw=1)
    ax.text(
        xs[0] + panel_w / 2,
        0.505,
        "Known-family + exact-identity controls\nLocked train / validation / test identities",
        ha="center",
        va="center",
        fontsize=7.1,
        color=MUTED,
        linespacing=1.35,
    )

    stage_header(ax, xs[1] + 0.035, 0.785, "2", "Paired Swin-S training", BLUE)
    rounded_box(ax, xs[1] + 0.025, 0.64, 0.098, 0.085, WHITE, BLUE, 0.01)
    ax.text(xs[1] + 0.074, 0.682, "Flat", ha="center", va="center", fontsize=8.8, fontweight="bold", color=BLUE)
    ax.text(xs[1] + 0.074, 0.648, "condition head", ha="center", va="center", fontsize=6.5, color=MUTED)
    rounded_box(ax, xs[1] + 0.158, 0.64, 0.102, 0.085, WHITE, BLUE, 0.01)
    ax.text(xs[1] + 0.209, 0.682, "Dual", ha="center", va="center", fontsize=8.8, fontweight="bold", color=BLUE)
    ax.text(xs[1] + 0.209, 0.648, "condition + host", ha="center", va="center", fontsize=6.5, color=MUTED)
    ax.text(
        xs[1] + panel_w / 2,
        0.578,
        "10 paired seeds per primary model",
        ha="center",
        va="center",
        fontsize=7.9,
        fontweight="bold",
        color=INK,
    )
    ax.text(
        xs[1] + panel_w / 2,
        0.523,
        "5 matched seeds: MAE and KL(q||p^h)\n30 epochs per run  |  90 runs total",
        ha="center",
        va="center",
        fontsize=7.2,
        color=MUTED,
        linespacing=1.35,
    )

    stage_header(ax, xs[2] + 0.035, 0.785, "3", "Locked-test audit", AMBER)
    metrics = (("Accuracy", BLUE_LIGHT, BLUE), ("Macro-F1", GREEN_LIGHT, GREEN), ("CHER", RED_LIGHT, RED))
    for index, (label, face, edge) in enumerate(metrics):
        x = xs[2] + 0.018 + index * 0.087
        rounded_box(ax, x, 0.658, 0.076, 0.058, face, edge, 0.009)
        ax.text(x + 0.038, 0.687, label, ha="center", va="center", fontsize=7.1, fontweight="bold", color=edge)
    ax.text(
        xs[2] + panel_w / 2,
        0.604,
        "Paired equivalence + HDR + TCI",
        ha="center",
        va="center",
        fontsize=7.8,
        fontweight="bold",
        color=INK,
    )
    ax.text(
        xs[2] + panel_w / 2,
        0.539,
        "Family balance  |  long tail  |  host-level error\nOracle host  |  calibration  |  selective risk",
        ha="center",
        va="center",
        fontsize=6.9,
        color=MUTED,
        linespacing=1.35,
    )

    arrow(ax, xs[0] + panel_w + 0.005, xs[1] - 0.007, 0.645)
    arrow(ax, xs[1] + panel_w + 0.005, xs[2] - 0.007, 0.645)

    rounded_box(ax, 0.035, 0.205, 0.45, 0.19, BLUE_LIGHT, BLUE)
    ax.text(0.055, 0.355, "Predicted host supervision", fontsize=9.0, fontweight="bold", color=BLUE, va="center")
    ax.text(
        0.055,
        0.293,
        "Equivalent on all 43-class primary outcomes",
        fontsize=7.9,
        color=INK,
        va="center",
    )
    ax.text(
        0.055,
        0.243,
        "TCI changed mean accuracy by at most 0.018 percentage points",
        fontsize=6.8,
        color=MUTED,
        va="center",
    )

    rounded_box(ax, 0.515, 0.205, 0.45, 0.19, GREEN_LIGHT, GREEN)
    ax.text(0.535, 0.355, "Ground-truth-host oracle", fontsize=9.0, fontweight="bold", color=GREEN, va="center")
    ax.text(
        0.535,
        0.293,
        "+2.2 points accuracy on the 215-class benchmark",
        fontsize=7.9,
        color=INK,
        va="center",
    )
    ax.text(
        0.535,
        0.243,
        "+10.2 to +10.6 points macro-F1, without refitting",
        fontsize=7.0,
        color=MUTED,
        va="center",
    )

    ax.add_patch(FancyBboxPatch((0, 0), 1, 0.145, boxstyle="square,pad=0", facecolor=INK, edgecolor=INK))
    ax.text(
        0.5,
        0.093,
        "Consistency is not correctness",
        ha="center",
        va="center",
        fontsize=11.5,
        fontweight="bold",
        color=WHITE,
    )
    ax.text(
        0.5,
        0.047,
        "Use hard host constraints only when the host information is independently reliable.",
        ha="center",
        va="center",
        fontsize=7.8,
        color="#DCE6E8",
    )

    for suffix in ("png", "pdf"):
        fig.savefig(
            OUTPUT_DIR / f"Graphical_Abstract.{suffix}",
            dpi=300,
            facecolor=WHITE,
            bbox_inches=None,
            pad_inches=0,
        )
    fig.savefig(
        OUTPUT_DIR / "Graphical_Abstract.tiff",
        dpi=300,
        facecolor=WHITE,
        bbox_inches=None,
        pad_inches=0,
        pil_kwargs={"compression": "tiff_lzw"},
    )
    plt.close(fig)


def make_highlights_docx() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    document = Document()
    document.core_properties.title = f"Highlights: {TITLE}"
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    heading = document.add_paragraph()
    run = heading.add_run("Highlights")
    run.bold = True
    run.font.size = Pt(14)
    document.add_paragraph(TITLE)
    for highlight in HIGHLIGHTS:
        if len(highlight) > 85:
            raise ValueError(f"Highlight exceeds 85 characters: {highlight}")
        document.add_paragraph(highlight, style="List Bullet")
    document.save(OUTPUT_DIR / "Highlights.docx")


def main() -> None:
    make_graphical_abstract()
    make_highlights_docx()
    for highlight in HIGHLIGHTS:
        print(f"{len(highlight):2d}  {highlight}")


if __name__ == "__main__":
    main()
